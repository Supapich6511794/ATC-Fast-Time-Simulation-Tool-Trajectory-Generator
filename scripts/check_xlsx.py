"""Check that a generated .xlsx really is one, without opening Excel.

The web app writes its chart workbooks by hand (``web/lib/report/xlsx.ts``) —
no library sits between us and the OOXML spec, so a mistake there produces a
file Excel refuses outright rather than one it renders imperfectly. This script
is the ground check for that:

    python scripts/check_xlsx.py path/to/file.xlsx

It always verifies the container: the ZIP index, every entry's CRC, that each
part is well-formed XML, and that the parts a charted workbook needs are all
present and wired to each other. If ``openpyxl`` happens to be installed it also
opens the workbook the way a spreadsheet application would and prints the sheet
and the chart it found, which is the check that actually proves the file works.

Exit status is 0 when everything passed, 1 otherwise, so it can be dropped into
a pipeline.
"""

from __future__ import annotations

import sys
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path

REQUIRED_PARTS = [
    "[Content_Types].xml",
    "_rels/.rels",
    "xl/workbook.xml",
    "xl/_rels/workbook.xml.rels",
    "xl/worksheets/sheet1.xml",
]

# Only needed when the workbook draws a chart, which is every one this app
# writes; checked separately so a plain data workbook is not failed for it.
CHART_PARTS = [
    "xl/chartsheets/sheet1.xml",
    "xl/chartsheets/_rels/sheet1.xml.rels",
    "xl/drawings/drawing1.xml",
    "xl/drawings/_rels/drawing1.xml.rels",
    "xl/charts/chart1.xml",
]


def check(path: Path) -> bool:
    ok = True
    print(f"== {path.name}  ({path.stat().st_size:,} bytes)")

    if not zipfile.is_zipfile(path):
        print("   NOT a ZIP archive — an .xlsx is a ZIP of XML parts")
        return False

    with zipfile.ZipFile(path) as z:
        bad = z.testzip()
        if bad is None:
            print("   CRC of every entry: OK")
        else:
            print(f"   CRC FAILED on {bad}")
            ok = False

        names = z.namelist()
        for part in REQUIRED_PARTS:
            if part not in names:
                print(f"   MISSING required part: {part}")
                ok = False

        missing_chart = [p for p in CHART_PARTS if p not in names]
        if missing_chart == CHART_PARTS:
            print("   no chart parts — this workbook holds data only")
        elif missing_chart:
            print(f"   chart is INCOMPLETE, missing: {', '.join(missing_chart)}")
            ok = False

        for name in names:
            if not name.endswith((".xml", ".rels")):
                continue
            try:
                ET.fromstring(z.read(name))
            except ET.ParseError as exc:
                print(f"   MALFORMED XML in {name}: {exc}")
                ok = False

        # The chain Excel follows: sheet -> drawing -> chart. A part that is
        # present but unreferenced is invisible, which looks like a chart that
        # silently did not appear.
        if "xl/charts/chart1.xml" in names:
            tab = z.read("xl/chartsheets/sheet1.xml").decode("utf-8")
            if "<drawing" not in tab:
                print("   chart tab does not reference its drawing — no chart will show")
                ok = False
            # A worksheet that names a drawing it has no relationship for is a
            # dangling reference, and Excel refuses the whole workbook for it.
            for part in (n for n in names if n.startswith("xl/worksheets/sheet")):
                if part.endswith(".rels"):
                    continue
                body = z.read(part).decode("utf-8")
                rels = part.replace("worksheets/", "worksheets/_rels/") + ".rels"
                if "<drawing" in body and rels not in names:
                    print(f"   {part} references a drawing with no relationship part")
                    ok = False
            chart = z.read("xl/charts/chart1.xml").decode("utf-8")
            refs = chart.count("<c:f>")
            print(f"   chart references {refs} cell range(s)")
            if refs == 0:
                print("   chart plots NOTHING — no ranges referenced")
                ok = False

    try:
        import openpyxl  # noqa: PLC0415  (optional, checked at runtime)
    except ImportError:
        print("   (install openpyxl for the full open-it-and-look check)")
        return ok

    import warnings

    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wb = openpyxl.load_workbook(path)
    # Every tab, not just the worksheets: the chart lives on a chartsheet, and
    # `wb.worksheets` skips those — a walk over it alone reports "no charts" on
    # a workbook whose whole point is the chart.
    found = 0
    for name in wb.sheetnames:
        tab = wb[name]
        charts = getattr(tab, "_charts", [])
        found += len(charts)
        if hasattr(tab, "dimensions"):
            print(f"   sheet {name!r}: {tab.dimensions}, {tab.max_row} rows")
        else:
            print(f"   chart tab {name!r}")
        for ch in charts:
            title = "titled" if ch.title is not None else "untitled"
            print(f"     {type(ch).__name__}, {len(ch.series)} series, {title}")
    if found == 0:
        print("   NO chart found by the reader — the file opens, but empty")
        ok = False
    return ok


def main(argv: list[str]) -> int:
    if len(argv) < 2:
        print(__doc__)
        return 1
    results = [check(Path(a)) for a in argv[1:]]
    print("\nall checks passed" if all(results) else "\nFAILED")
    return 0 if all(results) else 1


if __name__ == "__main__":
    raise SystemExit(main(sys.argv))
